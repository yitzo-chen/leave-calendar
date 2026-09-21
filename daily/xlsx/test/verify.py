# 驗證產生的 xlsx：python verify.py <xlsx路徑>
# 需要同目錄的 <xlsx路徑>.expected.json（由 build-sample.js 產生，記錄逐日累計的預期值）
import sys, zipfile, re, json, os
from lxml import etree
import openpyxl

TEMPLATE = r'D:\洲際液化天然氣接收站AI\日報範本.xlsx'
path = sys.argv[1]
ok = True


def check(cond, msg):
    global ok
    print(('PASS ' if cond else 'FAIL ') + msg)
    if not cond:
        ok = False


z = zipfile.ZipFile(path)
tz = zipfile.ZipFile(TEMPLATE)

# 1. 所有 XML 零件格式正確、沒有資料夾項目
bad = []
for n in [x for x in z.namelist() if not x.endswith('/')]:
    try:
        etree.fromstring(z.read(n))
    except Exception as e:
        bad.append((n, str(e)[:80]))
check(not bad and not any(x.endswith('/') for x in z.namelist()),
      f'所有 {len(z.namelist())} 個零件都是格式正確的 XML、且沒有多餘的資料夾項目 {bad or ""}')
check(z.namelist()[0] == '[Content_Types].xml', '[Content_Types].xml 為壓縮檔第一個項目')

wb = openpyxl.load_workbook(path)
tpl = openpyxl.load_workbook(TEMPLATE)['115.9.13']
names = wb.sheetnames
check(names == sorted(names, key=lambda s: tuple(map(int, s.split('.')))), f'分頁依日期排序 {names}')

# 2. 列印設定：每個分頁都要與範本相同
for n in names:
    ws = wb[n]
    check(ws.page_setup.paperSize == 9 and ws.page_setup.orientation == 'portrait' and ws.page_setup.scale == 94, f'[{n}] A4/直向/scale94')
    check(ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage, f'[{n}] fitToPage')
    check(ws.print_area == f"'{n}'!$A$1:$S$55", f'[{n}] 列印範圍 {ws.print_area}')
    check(ws.print_options.horizontalCentered and ws.print_options.verticalCentered, f'[{n}] 水平/垂直置中')
    pm, tm = ws.page_margins, tpl.page_margins
    check(abs(pm.left - tm.left) < 1e-9 and abs(pm.top - tm.top) < 1e-9 and abs(pm.header - tm.header) < 1e-9, f'[{n}] 邊界')
check(sum(1 for n in names if wb[n].sheet_view.tabSelected) == 1, '只有一個分頁被選取（避免 Excel 群組編輯）')


# 3. 儲存格：直接比對 XML 原文與樣式
def rng(a, b, cols):
    return {f'{c}{r}' for r in range(a, b + 1) for c in cols}


NAME_CELLS = rng(8, 18, 'A') | rng(20, 29, 'A') | rng(31, 45, 'A')
SHRINK_CELLS = rng(34, 45, 'MP') | {'J34', 'J37', 'J40', 'J43', 'O53'}
STYLE_CHANGED = NAME_CELLS | SHRINK_CELLS | {'A46', 'J5'}      # 預期樣式編號會變（衍生新樣式）
WRITTEN = (rng(8, 18, 'ADEF') | rng(20, 29, 'ADEF') | rng(31, 45, 'ADF') | rng(8, 20, 'H') | rng(22, 32, 'H')
           | rng(34, 45, 'MOP') | {'P33', 'A1', 'A3', 'A4', 'A5', 'A6', 'J5', 'N5', 'P5', 'J6', 'J34', 'J37', 'J40', 'J43', 'L34', 'L40', 'A46', 'O53'})

CELL = re.compile(r'<c r="([A-Z]+\d+)"((?:\s+[A-Za-z:]+="[^"]*")*)\s*(?:/>|>[\s\S]*?</c>)')


def cells_of(xml):
    out = {}
    for m in CELL.finditer(xml):
        sm = re.search(r'\ss="(\d+)"', m.group(2))
        out[m.group(1)] = (m.group(0), int(sm.group(1)) if sm else None)
    return out


def xfs_of(zf):
    st = zf.read('xl/styles.xml').decode('utf-8')
    body = re.search(r'<cellXfs count="\d+">(.*?)</cellXfs>', st, re.S).group(1)
    return re.findall(r'<xf [^>]*?(?:/>|>.*?</xf>)', body, re.S), st


def attr(x, name):
    m = re.search(name + r'="([^"]*)"', x)
    return m.group(1) if m else None


txfs, tst = xfs_of(tz)
oxfs, ost = xfs_of(z)
check(oxfs[:len(txfs)] == txfs, f'既有 {len(txfs)} 個樣式原封不動，新樣式只是附加在後面（共 {len(oxfs)} 個）')


def fonts_of(st):
    body = re.search(r'<fonts[^>]*>(.*?)</fonts>', st, re.S).group(1)
    return re.findall(r'<font>.*?</font>|<font/>', body, re.S)


ofonts = fonts_of(ost)
tcells = cells_of(tz.read('xl/worksheets/sheet1.xml').decode('utf-8'))
wbxml = z.read('xl/workbook.xml').decode('utf-8')
rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
sheet_xml = {}
for n in names:
    rid = re.search(r'<sheet name="%s"[^>]*r:id="([^"]+)"' % re.escape(n), wbxml).group(1)
    target = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid, rels).group(1)
    sheet_xml[n] = z.read('xl/' + target).decode('utf-8')

for n in names:
    ocells = cells_of(sheet_xml[n])
    check(set(ocells) == set(tcells), f'[{n}] 儲存格集合與範本相同（{len(tcells)} 格）')
    untouched = [k for k in tcells if k not in WRITTEN]
    diff_raw = [k for k in untouched if ocells[k][0] != tcells[k][0]]
    check(not diff_raw, f'[{n}] 未動儲存格 XML 原文與範本完全相同（{len(untouched)} 格） {diff_raw[:8]}')
    same_style = [k for k in WRITTEN - STYLE_CHANGED if ocells[k][1] != tcells[k][1]]
    check(not same_style, f'[{n}] 已寫入但不換樣式的儲存格，樣式編號不變 {same_style[:8]}')
    # 換樣式的儲存格：框線/填色/數字格式必須與範本相同，只有字型與對齊可變
    bad_border = []
    for k in STYLE_CHANGED:
        tx, ox = txfs[tcells[k][1]], oxfs[ocells[k][1]]
        if any(attr(tx, a) != attr(ox, a) for a in ('borderId', 'fillId', 'numFmtId')):
            bad_border.append(k)
    check(not bad_border, f'[{n}] 換樣式的 {len(STYLE_CHANGED)} 格，框線/填色/數字格式與範本相同 {bad_border[:8]}')
    ws = wb[n]
    check(all(ws[k].alignment.shrink_to_fit and not ws[k].alignment.wrap_text for k in (NAME_CELLS | SHRINK_CELLS)),
          f'[{n}] 名稱/人員/加班原因/填表人格皆為「縮小字型以符合儲存格」')
    sizes = {ws[k].font.sz for k in NAME_CELLS}
    check(sizes == {12.0}, f'[{n}] 工種/機具/材料名稱格字型大小一致 {sizes}（範本原為 8pt 與 12pt 混用）')
    check((ws['A46'].alignment.horizontal, ws['A46'].alignment.vertical, ws['A46'].alignment.wrap_text) == ('left', 'top', True), f'[{n}] A46 備註方塊 向左靠齊+靠上+自動換行')
    check([ws.column_dimensions[k].width for k in 'ABCDEFGHIJKLMNOPQRS'] == [tpl.column_dimensions[k].width for k in 'ABCDEFGHIJKLMNOPQRS'], f'[{n}] 欄寬')
    check([ws.row_dimensions[r].height for r in range(1, 58)] == [tpl.row_dimensions[r].height for r in range(1, 58)], f'[{n}] 列高')
    check(re.search(r'<pageSetup[^>]*paperSize="9"[^>]*scale="94"', sheet_xml[n]) is not None, f'[{n}] pageSetup XML 保留 paperSize/scale')

# 4. 合併儲存格
ws = wb[names[-1]]
ms = {str(m) for m in ws.merged_cells.ranges}
tms = {str(m) for m in tpl.merged_cells.ranges}
in_note = lambda m: (lambda r: 46 <= r[1] <= 52 and r[0] >= 1 and r[2] <= 19)(__import__('openpyxl').utils.range_boundaries(m))
expected_ms = {m for m in tms if not in_note(m)} | {'A46:S52', 'J5:L5'} | {f'P{r}:S{r}' for r in range(33, 46)}
check(ms == expected_ms, f'合併數 {len(ms)}（預期 = 範本非備註區合併 + A46:S52 + J5:L5 + 13 個 P:S = {len(expected_ms)}）差異 {sorted(ms ^ expected_ms)[:6]}')
check('A46:S52' in ms and 'A46:B46' not in ms and 'Q51:R51' not in ms, '備註區併成 A46:S52，原小合併已移除')
check(all(f'P{r}:S{r}' in ms for r in range(33, 46)), '加班原因欄 P:S 逐列合併（33–45）')
check({'H34:H39', 'L34:L39', 'J34:K36', 'O53:S54'} <= ms, '其他範本合併保留')

# 5. 逐日累計：每個分頁的累計 = 「當天＋之前」，不是同一個最終值
exp = json.load(open(path + '.expected.json', encoding='utf-8')) if os.path.exists(path + '.expected.json') else {}
for n in names:
    w = wb[n]
    bad = {k: (w[k].value, v) for k, v in exp.get(n, {}).items() if w[k].value != v}
    check(exp.get(n) and not bad, f'[{n}] 名稱與累計欄 {len(exp.get(n, {}))} 格與預期一致 {bad}')
series = [wb[n]['F8'].value for n in names]
check(series == sorted(series) and len(set(series)) == len(series), f'公司工累計逐日遞增 {dict(zip(names, series))}')
carry = [wb[n]['F9'].value for n in names]
check(carry[1] == carry[0], f'模板工 9/19 沒出工，累計延續前一天 {dict(zip(names, carry))}')

# 6. 欄位內容抽查（最後一個分頁）
ws = wb[names[-1]]


def val(k):
    return ws[k].value


LONG = '樓梯防滑地磚(20*27)含黏著劑及施工損耗合計(箱)'
exp_cells = {
    'A1': None, 'A3': '業主：中油', 'A4': '工程名稱：洲際液化天然氣接收站', 'A5': '合約金額：', 'A6': '開工日期：',
    'J5': '日期：115', 'N5': 9, 'P5': 20,
    'J6': '    天氣：  □晴 □陰   ■雨   施工狀況：■施工  □休息',
    'A8': '公司工', 'D8': 4, 'E8': 4, 'A9': '模板工', 'D9': 1, 'E9': 1, 'A10': '鋼筋工(工地)', 'D10': None, 'A12': None, 'F12': None,
    'A20': '吊車', 'D20': 2, 'E20': 2,
    'A31': '210kg/cm²混凝土(m³)', 'D31': None, 'A33': LONG, 'D33': 1,
    'J34': '王大明', 'J37': '李小華', 'J40': '王大明', 'J43': '李小華', 'L34': '本工', 'L40': '本工',
    'M34': '陳一', 'M35': '林二', 'M36': None, 'O35': None,
    'M40': '陳一', 'M41': '林二', 'O41': 3, 'P41': '配合泵送車進場時間安排夜間趕工',
    'P33': '加班原因', 'O33': '加班', 'M33': '人員', 'A46': '備註：第一行備註\n第二行備註', 'O53': '陳易佐', 'M53': '填表人',
}
bad = {k: (val(k), v) for k, v in exp_cells.items() if val(k) != v}
check(not bad, f'[{names[-1]}] 欄位內容 {len(exp_cells)} 項抽查（含超長材料名稱完整保留）{bad}')
lines = [val(f'H{r}') for r in range(8, 21) if val(f'H{r}')]
check(''.join(lines[1:]).startswith('2.二樓鋼筋綁紮') and '調度' in ''.join(lines), f'長文字自動折行 {len(lines)} 行')
print('\nRESULT:', 'ALL PASS' if ok else 'HAS FAIL')
sys.exit(0 if ok else 1)
